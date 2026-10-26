"""Excel and PDF export smoke tests (open generated files)."""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from app.models import Correction, Document, Sheet
from app.services.exports import (
    SheetExportBundle,
    build_document_quote_pdf,
    build_document_workbook,
    build_sheet_quote_pdf,
    build_sheet_workbook,
)
from app.services.takeoff_svc import sheet_bom
from app.settings import REPO_ROOT, Settings, get_settings
from openpyxl import load_workbook
from takeoff_engine.models import SheetResult
from takeoff_engine.takeoff import compute_takeoff


def _bundle_from_fixture() -> SheetExportBundle:
    path = REPO_ROOT / "contracts" / "fixtures" / "sheet_result.sample.json"
    result = SheetResult.model_validate(json.loads(path.read_text(encoding="utf-8")))
    doc = Document(
        id="doc-fixture",
        filename="escondido_plans.pdf",
        file_hash="a" * 64,
        storage_path="storage/documents/fixture.pdf",
        page_count=4,
    )
    sheet = Sheet(
        id=result.sheet.id,
        document_id=doc.id,
        page_index=0,
        sheet_number=result.sheet.sheet_number,
        title=result.sheet.title,
        page_width=result.sheet.page_width,
        page_height=result.sheet.page_height,
        viewport=result.sheet.viewport.model_dump(),
        scale_pt_per_ft=result.sheet.scale_pt_per_ft,
    )
    sheet.document = doc
    takeoff = compute_takeoff(result.rooms, result.detections)
    settings = get_settings()
    bom = sheet_bom(settings, sheet.id, result.detections, result.wire_runs, result.boxes)
    return SheetExportBundle(
        sheet=sheet,
        document=doc,
        rooms=result.rooms,
        detections=result.detections,
        wires=result.wire_runs,
        takeoff=takeoff,
        bom=bom,
        corrections=[
            Correction(
                id="corr-1",
                sheet_id=sheet.id,
                analysis_run_id="run-1",
                entity="detection",
                entity_id="det-1",
                action="update",
                before={"x": 1},
                after={"x": 2},
            )
        ],
        pipeline_version="0.1.0",
        overlay_png=None,
    )


def test_sheet_workbook_has_expected_tabs_and_content() -> None:
    bundle = _bundle_from_fixture()
    raw = build_sheet_workbook(bundle)
    wb = load_workbook(BytesIO(raw))
    names = set(wb.sheetnames)
    assert names == {
        "Takeoff by room",
        "Detections",
        "BOM",
        "Wires",
        "Assumptions",
        "Corrections log",
    }
    takeoff = wb["Takeoff by room"]
    assert takeoff["A1"].value == "Room"
    assert takeoff.freeze_panes == "A2"
    assert takeoff.max_row >= 2

    detections = wb["Detections"]
    headers = [detections.cell(1, c).value for c in range(1, 11)]
    assert headers == [
        "id",
        "category",
        "subtype",
        "codes",
        "room",
        "x",
        "y",
        "confidence",
        "source",
        "needs_review",
    ]
    assert detections.max_row >= 2

    bom = wb["BOM"]
    assert bom["A1"].value == "item_key"
    assert bom.max_row >= 2

    corr = wb["Corrections log"]
    assert corr["A2"].value == "corr-1"


def test_document_workbook_has_plan_column() -> None:
    b1 = _bundle_from_fixture()
    b2 = _bundle_from_fixture()
    b2.sheet.page_index = 1
    b2.sheet.sheet_number = "A2-111"
    raw = build_document_workbook([b1, b2])
    wb = load_workbook(BytesIO(raw))
    assert "Takeoff by room" in wb.sheetnames
    takeoff = wb["Takeoff by room"]
    assert takeoff["A1"].value == "Plan"
    assert takeoff.max_row >= 3


def test_sheet_quote_pdf_is_valid_pdf() -> None:
    bundle = _bundle_from_fixture()
    settings = Settings.model_construct(
        storage_dir=REPO_ROOT / "storage",
        cache_dir=REPO_ROOT / ".cache",
        enable_bom=True,
        price_list_path=REPO_ROOT / "data" / "price_list.yaml",
        assumptions_path=REPO_ROOT / "data" / "assumptions.yaml",
    )
    raw = build_sheet_quote_pdf(bundle, settings)
    assert raw.startswith(b"%PDF")
    assert b"ReportLab" in raw
    assert len(raw) > 800


def test_document_quote_pdf_multipage() -> None:
    b1 = _bundle_from_fixture()
    b2 = _bundle_from_fixture()
    b2.sheet.page_index = 1
    b2.sheet.sheet_number = "A2-111"
    settings = get_settings()
    raw = build_document_quote_pdf([b1, b2], settings)
    assert raw.startswith(b"%PDF")
    assert raw.count(b"/Type /Page") >= 2 or raw.count(b"/Page") >= 2


@pytest.mark.asyncio
async def test_export_endpoints(client, sample_pdf: Path) -> None:
    with sample_pdf.open("rb") as f:
        up = await client.post(
            "/api/documents",
            files={"file": ("escondido_plans.pdf", f, "application/pdf")},
        )
    assert up.status_code == 202, up.text
    doc_id = up.json()["document_id"]
    sheet_id = up.json()["sheets"][0]["sheet_id"]

    xlsx = await client.get(f"/api/sheets/{sheet_id}/export.xlsx")
    assert xlsx.status_code == 200, xlsx.text
    assert xlsx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = load_workbook(BytesIO(xlsx.content))
    assert "Takeoff by room" in wb.sheetnames
    assert "Detections" in wb.sheetnames

    pdf = await client.get(f"/api/sheets/{sheet_id}/quote.pdf")
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")

    dx = await client.get(f"/api/documents/{doc_id}/export.xlsx")
    assert dx.status_code == 200
    dwb = load_workbook(BytesIO(dx.content))
    assert dwb["Takeoff by room"]["A1"].value == "Plan"

    dp = await client.get(f"/api/documents/{doc_id}/quote.pdf")
    assert dp.status_code == 200
    assert dp.content.startswith(b"%PDF")
