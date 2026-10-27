"""Cable / wire-run edit API (isolated takeoff_test DB)."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook


async def _upload(client: AsyncClient, sample_pdf: Path) -> dict:
    with sample_pdf.open("rb") as f:
        resp = await client.post(
            "/api/documents",
            files={"file": ("escondido_plans.pdf", f, "application/pdf")},
        )
    assert resp.status_code == 202, resp.text
    return resp.json()


async def _two_active_dets(client: AsyncClient, sheet_id: str) -> tuple[str, str]:
    sheet = await client.get(f"/api/sheets/{sheet_id}")
    assert sheet.status_code == 200
    active = [d for d in sheet.json()["detections"] if d["status"] == "active"]
    assert len(active) >= 2
    return active[0]["id"], active[1]["id"]


@pytest.mark.asyncio
async def test_list_create_patch_delete_restore_wire(
    client: AsyncClient, sample_pdf: Path
) -> None:
    up = await _upload(client, sample_pdf)
    sheet_id = up["sheets"][0]["sheet_id"]
    a, b = await _two_active_dets(client, sheet_id)

    listed = await client.get(f"/api/sheets/{sheet_id}/wire-runs")
    assert listed.status_code == 200
    before_n = len(listed.json())

    created = await client.post(
        f"/api/sheets/{sheet_id}/wire-runs",
        json={"from_detection_id": a, "to_detection_id": b, "cable_type": "nm_14_3"},
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["wire_run"]["cable_type"] == "nm_14_3"
    assert body["wire_run"]["source"] == "user"
    assert body["bom"] is not None
    wire_id = body["wire_run"]["id"]
    assert len(body["wire_run"]["points"]) >= 2

    listed2 = await client.get(f"/api/sheets/{sheet_id}/wire-runs")
    assert len(listed2.json()) == before_n + 1

    # Manual cable_type sticks after points reshape (source=user)
    patched = await client.patch(
        f"/api/wire-runs/{wire_id}",
        json={
            "points": [
                {"x": 10.0, "y": 10.0},
                {"x": 110.0, "y": 10.0},
            ]
        },
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["wire_run"]["cable_type"] == "nm_14_3"
    assert patched.json()["wire_run"]["plan_length_ft"] > 0

    deleted = await client.delete(f"/api/wire-runs/{wire_id}")
    assert deleted.status_code == 200
    assert deleted.json()["bom"] is not None
    listed3 = await client.get(f"/api/sheets/{sheet_id}/wire-runs")
    assert wire_id not in {w["id"] for w in listed3.json()}

    restored = await client.post(f"/api/wire-runs/{wire_id}/restore")
    assert restored.status_code == 200, restored.text
    assert restored.json()["wire_run"]["id"] == wire_id
    listed4 = await client.get(f"/api/sheets/{sheet_id}/wire-runs")
    assert wire_id in {w["id"] for w in listed4.json()}

    corr = await client.get(f"/api/sheets/{sheet_id}/corrections")
    actions = [(c["entity"], c["action"]) for c in corr.json()]
    assert ("wire_run", "create") in actions
    assert ("wire_run", "delete") in actions
    assert ("wire_run", "restore") in actions


@pytest.mark.asyncio
async def test_move_detection_reattaches_runs(
    client: AsyncClient, sample_pdf: Path
) -> None:
    up = await _upload(client, sample_pdf)
    sheet_id = up["sheets"][0]["sheet_id"]
    a, b = await _two_active_dets(client, sheet_id)

    created = await client.post(
        f"/api/sheets/{sheet_id}/wire-runs",
        json={"from_detection_id": a, "to_detection_id": b},
    )
    assert created.status_code == 201
    wire_id = created.json()["wire_run"]["id"]
    old_pts = created.json()["wire_run"]["points"]

    sheet = await client.get(f"/api/sheets/{sheet_id}")
    det = next(d for d in sheet.json()["detections"] if d["id"] == a)
    new_center = {"x": det["center"]["x"] + 40.0, "y": det["center"]["y"] + 25.0}

    moved = await client.patch(
        f"/api/detections/{a}",
        json={"center": new_center},
    )
    assert moved.status_code == 200, moved.text
    payload = moved.json()
    assert payload["detection"]["center"]["x"] == pytest.approx(new_center["x"])
    assert payload["bom"] is not None
    assert payload["boxes"] is not None
    assert isinstance(payload["boxes"], list)
    changed = payload.get("wire_runs") or []
    assert any(w["id"] == wire_id for w in changed)
    new_wire = next(w for w in changed if w["id"] == wire_id)
    assert new_wire["points"][0]["x"] == pytest.approx(new_center["x"])
    assert new_wire["points"][0] != old_pts[0]


def _box_gang_signature(boxes: list[dict]) -> list[tuple[int, int]]:
    """Stable (gangs, n_detections) multiset for comparing box lists."""
    return sorted((int(b["gangs"]), len(b.get("detection_ids") or [])) for b in boxes)


@pytest.mark.asyncio
async def test_move_patch_response_matches_subsequent_get(
    client: AsyncClient, sample_pdf: Path
) -> None:
    """PATCH boxes/bom must equal a follow-up GET (no stale identity-map)."""
    up = await _upload(client, sample_pdf)
    sheet_id = up["sheets"][0]["sheet_id"]

    before = await client.get(f"/api/sheets/{sheet_id}")
    assert before.status_code == 200
    two_gang = next(
        (b for b in before.json()["boxes"] if int(b.get("gangs") or 1) >= 2),
        None,
    )
    assert two_gang is not None, "fixture should include a 2-gang switch box"
    det_id = two_gang["detection_ids"][0]
    det = next(d for d in before.json()["detections"] if d["id"] == det_id)
    # Move far enough to break the gang group.
    new_center = {"x": det["center"]["x"] + 80.0, "y": det["center"]["y"] + 80.0}

    moved = await client.patch(
        f"/api/detections/{det_id}",
        json={"center": new_center},
    )
    assert moved.status_code == 200, moved.text
    payload = moved.json()
    assert payload["boxes"] is not None
    assert payload["bom"] is not None

    after = await client.get(f"/api/sheets/{sheet_id}")
    assert after.status_code == 200
    sheet_json = after.json()
    bom_get = await client.get(f"/api/sheets/{sheet_id}/bom")
    assert bom_get.status_code == 200

    assert _box_gang_signature(payload["boxes"]) == _box_gang_signature(
        sheet_json["boxes"]
    )
    # Breaking a 2G should yield no remaining 2G that still lists this det.
    assert not any(
        int(b.get("gangs") or 1) >= 2 and det_id in (b.get("detection_ids") or [])
        for b in payload["boxes"]
    )
    assert payload["bom"]["total_usd"] == pytest.approx(
        bom_get.json()["total_usd"], rel=1e-6, abs=1e-6
    )
    # Box / plate lines should match GET bom.
    def _box_bom_lines(bom: dict) -> set[tuple[str, float]]:
        return {
            (ln["item_key"], float(ln["quantity"]))
            for ln in bom["lines"]
            if ln["item_key"].startswith("box_") or ln["item_key"].startswith("plate_")
        }

    assert _box_bom_lines(payload["bom"]) == _box_bom_lines(bom_get.json())


@pytest.mark.asyncio
async def test_delete_detection_hides_runs_from_bom_undo_restores(
    client: AsyncClient, sample_pdf: Path
) -> None:
    up = await _upload(client, sample_pdf)
    sheet_id = up["sheets"][0]["sheet_id"]
    a, b = await _two_active_dets(client, sheet_id)

    created = await client.post(
        f"/api/sheets/{sheet_id}/wire-runs",
        json={
            "from_detection_id": a,
            "to_detection_id": b,
            "cable_type": "nm_14_2",
        },
    )
    wire_id = created.json()["wire_run"]["id"]
    bom_before = created.json()["bom"]
    ft_before = sum(
        float(line["quantity"])
        for line in bom_before["lines"]
        if line["item_key"].endswith("_ft") and wire_id in line.get("source_ids", [])
    )

    deleted = await client.delete(f"/api/detections/{a}")
    assert deleted.status_code == 200
    assert deleted.json()["bom"] is not None
    wires = await client.get(f"/api/sheets/{sheet_id}/wire-runs")
    assert wire_id not in {w["id"] for w in wires.json()}

    # Soft-deleted detection: run still in DB — restore brings cable back
    restored = await client.post(f"/api/detections/{a}/restore")
    assert restored.status_code == 200
    wires2 = await client.get(f"/api/sheets/{sheet_id}/wire-runs")
    assert wire_id in {w["id"] for w in wires2.json()}
    bom_after = restored.json()["bom"]
    assert bom_after is not None
    if ft_before > 0:
        ft_after = sum(
            float(line["quantity"])
            for line in bom_after["lines"]
            if line["item_key"].endswith("_ft") and wire_id in line.get("source_ids", [])
        )
        assert ft_after > 0


@pytest.mark.asyncio
async def test_wire_undo_redo(client: AsyncClient, sample_pdf: Path) -> None:
    up = await _upload(client, sample_pdf)
    sheet_id = up["sheets"][0]["sheet_id"]
    a, b = await _two_active_dets(client, sheet_id)

    created = await client.post(
        f"/api/sheets/{sheet_id}/wire-runs",
        json={"from_detection_id": a, "to_detection_id": b},
    )
    wire_id = created.json()["wire_run"]["id"]

    hist = await client.get(f"/api/sheets/{sheet_id}/history")
    assert hist.json()["can_undo"] is True

    undone = await client.post(f"/api/sheets/{sheet_id}/undo")
    assert undone.status_code == 200
    wires = await client.get(f"/api/sheets/{sheet_id}/wire-runs")
    assert wire_id not in {w["id"] for w in wires.json()}

    redone = await client.post(f"/api/sheets/{sheet_id}/redo")
    assert redone.status_code == 200
    wires2 = await client.get(f"/api/sheets/{sheet_id}/wire-runs")
    assert wire_id in {w["id"] for w in wires2.json()}


@pytest.mark.asyncio
async def test_export_xlsx_reflects_edited_wire(
    client: AsyncClient, sample_pdf: Path
) -> None:
    up = await _upload(client, sample_pdf)
    sheet_id = up["sheets"][0]["sheet_id"]
    a, b = await _two_active_dets(client, sheet_id)

    created = await client.post(
        f"/api/sheets/{sheet_id}/wire-runs",
        json={
            "from_detection_id": a,
            "to_detection_id": b,
            "cable_type": "nm_14_3",
        },
    )
    wire_id = created.json()["wire_run"]["id"]

    xlsx = await client.get(f"/api/sheets/{sheet_id}/export.xlsx")
    assert xlsx.status_code == 200
    wb = load_workbook(BytesIO(xlsx.content))
    wires_sheet = wb["Wires"]
    ids = {
        wires_sheet.cell(r, 1).value
        for r in range(2, wires_sheet.max_row + 1)
    }
    assert wire_id in ids

    quote = await client.get(f"/api/sheets/{sheet_id}/quote.pdf")
    assert quote.status_code == 200
    assert quote.content[:4] == b"%PDF"
