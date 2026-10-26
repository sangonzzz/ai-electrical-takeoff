"""Excel takeoff workbooks and client-facing PDF quotes."""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from takeoff_engine.models import Bom, Detection, Room, Takeoff, WireRun

from app.models import Correction, Document, Sheet
from app.settings import Settings


@dataclass
class SheetExportBundle:
    sheet: Sheet
    document: Document
    rooms: list[Room]
    detections: list[Detection]
    wires: list[WireRun]
    takeoff: Takeoff
    bom: Bom
    corrections: list[Correction]
    pipeline_version: str
    overlay_png: Path | None = None


_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)
_CURRENCY = '"$"#,##0.00'


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = _THIN
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _autosize(ws, min_width: int = 10, max_width: int = 42) -> None:
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        length = 0
        for cell in col_cells:
            val = "" if cell.value is None else str(cell.value)
            length = max(length, len(val))
        ws.column_dimensions[letter].width = max(min_width, min(max_width, length + 2))


def _write_rows(
    ws,
    headers: list[str],
    rows: list[list[Any]],
    currency_cols: set[int] | None = None,
) -> None:
    currency_cols = currency_cols or set()
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.border = _THIN
            if c in currency_cols and isinstance(cell.value, (int, float)):
                cell.number_format = _CURRENCY
    _style_header(ws, len(headers))
    _autosize(ws)


def _room_name(rooms: list[Room], room_id: str | None) -> str:
    if not room_id:
        return ""
    for r in rooms:
        if r.id == room_id:
            return r.display_name
    return room_id


def _takeoff_type_keys(takeoff: Takeoff) -> list[str]:
    keys: set[str] = set()
    for row in takeoff.rows:
        keys.update(row.counts)
    keys.update(takeoff.totals)
    keys.update(takeoff.unassigned)
    # Prefer category keys first, then category.subtype
    cats = sorted(k for k in keys if "." not in k)
    subs = sorted(k for k in keys if "." in k)
    return cats + subs


def _append_takeoff_sheet(
    wb: Workbook, title: str, takeoff: Takeoff, *, plan_label: str | None = None
) -> None:
    ws = wb.create_sheet(title[:31])
    type_keys = _takeoff_type_keys(takeoff)
    headers = (["Plan"] if plan_label is not None else []) + ["Room"] + type_keys
    rows: list[list[Any]] = []
    for row in takeoff.rows:
        base = [plan_label] if plan_label is not None else []
        rows.append(base + [row.room_name] + [row.counts.get(k, 0) for k in type_keys])
    if takeoff.unassigned:
        base = [plan_label] if plan_label is not None else []
        rows.append(base + ["Unassigned"] + [takeoff.unassigned.get(k, 0) for k in type_keys])
    base = [plan_label] if plan_label is not None else []
    rows.append(base + ["TOTAL"] + [takeoff.totals.get(k, 0) for k in type_keys])
    _write_rows(ws, headers, rows)


def _append_detections_sheet(
    wb: Workbook,
    title: str,
    rooms: list[Room],
    detections: list[Detection],
    *,
    plan_label: str | None = None,
) -> None:
    ws = wb.create_sheet(title[:31])
    headers = (["Plan"] if plan_label is not None else []) + [
        "id",
        "category",
        "subtype",
        "codes",
        "room",
        "x",
        "y",
        "confidence",
        "source",
        "needs_review",
    ]
    rows: list[list[Any]] = []
    for d in detections:
        if d.status != "active":
            continue
        base = [plan_label] if plan_label is not None else []
        rows.append(
            base
            + [
                d.id,
                d.category.value,
                d.subtype,
                ", ".join(d.codes),
                _room_name(rooms, d.room_id),
                d.center.x,
                d.center.y,
                d.confidence,
                d.source,
                d.needs_review,
            ]
        )
    _write_rows(ws, headers, rows)


def _append_bom_sheet(
    wb: Workbook, title: str, bom: Bom, *, plan_label: str | None = None
) -> None:
    ws = wb.create_sheet(title[:31])
    headers = (["Plan"] if plan_label is not None else []) + [
        "item_key",
        "description",
        "quantity",
        "unit",
        "unit_price",
        "extended_price",
        "price_source",
    ]
    rows: list[list[Any]] = []
    for line in bom.lines:
        base = [plan_label] if plan_label is not None else []
        rows.append(
            base
            + [
                line.item_key,
                line.description,
                line.quantity,
                line.unit,
                line.unit_price,
                line.extended_price,
                line.price_source,
            ]
        )
    currency_offset = 1 if plan_label is not None else 0
    _write_rows(
        ws,
        headers,
        rows,
        currency_cols={5 + currency_offset, 6 + currency_offset},
    )


def _append_wires_sheet(
    wb: Workbook, title: str, wires: list[WireRun], *, plan_label: str | None = None
) -> None:
    ws = wb.create_sheet(title[:31])
    headers = (["Plan"] if plan_label is not None else []) + [
        "id",
        "cable_type",
        "plan_length_ft",
        "vertical_allowance_ft",
        "total_length_ft",
        "from_detection_id",
        "to_detection_id",
    ]
    rows: list[list[Any]] = []
    for w in wires:
        base = [plan_label] if plan_label is not None else []
        rows.append(
            base
            + [
                w.id,
                w.cable_type,
                w.plan_length_ft,
                w.vertical_allowance_ft,
                w.total_length_ft,
                w.from_detection_id,
                w.to_detection_id,
            ]
        )
    _write_rows(ws, headers, rows)


def _append_assumptions_sheet(
    wb: Workbook, title: str, bom: Bom, *, plan_label: str | None = None
) -> None:
    ws = wb.create_sheet(title[:31])
    headers = (["Plan"] if plan_label is not None else []) + ["key", "value"]
    rows: list[list[Any]] = []
    for key, value in sorted(bom.assumptions.items(), key=lambda kv: str(kv[0])):
        base = [plan_label] if plan_label is not None else []
        rows.append(base + [str(key), value])
    _write_rows(ws, headers, rows)


def _append_corrections_sheet(
    wb: Workbook,
    title: str,
    corrections: list[Correction],
    *,
    plan_label: str | None = None,
) -> None:
    ws = wb.create_sheet(title[:31])
    headers = (["Plan"] if plan_label is not None else []) + [
        "id",
        "entity",
        "entity_id",
        "action",
        "created_at",
        "before",
        "after",
    ]
    rows: list[list[Any]] = []
    for c in corrections:
        base = [plan_label] if plan_label is not None else []
        rows.append(
            base
            + [
                c.id,
                c.entity,
                c.entity_id,
                c.action,
                c.created_at.isoformat() if c.created_at else "",
                str(c.before) if c.before is not None else "",
                str(c.after) if c.after is not None else "",
            ]
        )
    _write_rows(ws, headers, rows)


def _plan_label(bundle: SheetExportBundle) -> str:
    s = bundle.sheet
    return s.sheet_number or s.title or f"page-{s.page_index + 1}"


def build_sheet_workbook(bundle: SheetExportBundle) -> bytes:
    wb = Workbook()
    # Remove default sheet
    default = wb.active
    wb.remove(default)
    _append_takeoff_sheet(wb, "Takeoff by room", bundle.takeoff)
    _append_detections_sheet(wb, "Detections", bundle.rooms, bundle.detections)
    _append_bom_sheet(wb, "BOM", bundle.bom)
    _append_wires_sheet(wb, "Wires", bundle.wires)
    _append_assumptions_sheet(wb, "Assumptions", bundle.bom)
    _append_corrections_sheet(wb, "Corrections log", bundle.corrections)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_document_workbook(bundles: list[SheetExportBundle]) -> bytes:
    """One workbook for all sheets; each tab includes a Plan column."""
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    # Combined tabs
    takeoff_ws = wb.create_sheet("Takeoff by room")
    det_ws = wb.create_sheet("Detections")
    bom_ws = wb.create_sheet("BOM")
    wires_ws = wb.create_sheet("Wires")
    assum_ws = wb.create_sheet("Assumptions")
    corr_ws = wb.create_sheet("Corrections log")

    # Build via temporary single-sheet helpers into memory is awkward; fill combined.
    all_type_keys: list[str] = []
    seen: set[str] = set()
    for b in bundles:
        for k in _takeoff_type_keys(b.takeoff):
            if k not in seen:
                seen.add(k)
                all_type_keys.append(k)

    takeoff_headers = ["Plan", "Room"] + all_type_keys
    takeoff_rows: list[list[Any]] = []
    for b in bundles:
        label = _plan_label(b)
        for row in b.takeoff.rows:
            takeoff_rows.append(
                [label, row.room_name] + [row.counts.get(k, 0) for k in all_type_keys]
            )
        if b.takeoff.unassigned:
            takeoff_rows.append(
                [label, "Unassigned"]
                + [b.takeoff.unassigned.get(k, 0) for k in all_type_keys]
            )
        takeoff_rows.append(
            [label, "TOTAL"] + [b.takeoff.totals.get(k, 0) for k in all_type_keys]
        )
    _write_rows(takeoff_ws, takeoff_headers, takeoff_rows)

    det_headers = [
        "Plan",
        "id",
        "category",
        "subtype",
        "codes",
        "room",
        "x",
        "y",
        "confidence",
        "source",
        "needs_review",
    ]
    det_rows: list[list[Any]] = []
    for b in bundles:
        label = _plan_label(b)
        for d in b.detections:
            if d.status != "active":
                continue
            det_rows.append(
                [
                    label,
                    d.id,
                    d.category.value,
                    d.subtype,
                    ", ".join(d.codes),
                    _room_name(b.rooms, d.room_id),
                    d.center.x,
                    d.center.y,
                    d.confidence,
                    d.source,
                    d.needs_review,
                ]
            )
    _write_rows(det_ws, det_headers, det_rows)

    bom_headers = [
        "Plan",
        "item_key",
        "description",
        "quantity",
        "unit",
        "unit_price",
        "extended_price",
        "price_source",
    ]
    bom_rows: list[list[Any]] = []
    for b in bundles:
        label = _plan_label(b)
        for line in b.bom.lines:
            bom_rows.append(
                [
                    label,
                    line.item_key,
                    line.description,
                    line.quantity,
                    line.unit,
                    line.unit_price,
                    line.extended_price,
                    line.price_source,
                ]
            )
    _write_rows(bom_ws, bom_headers, bom_rows, currency_cols={6, 7})

    wires_headers = [
        "Plan",
        "id",
        "cable_type",
        "plan_length_ft",
        "vertical_allowance_ft",
        "total_length_ft",
        "from_detection_id",
        "to_detection_id",
    ]
    wires_rows: list[list[Any]] = []
    for b in bundles:
        label = _plan_label(b)
        for w in b.wires:
            wires_rows.append(
                [
                    label,
                    w.id,
                    w.cable_type,
                    w.plan_length_ft,
                    w.vertical_allowance_ft,
                    w.total_length_ft,
                    w.from_detection_id,
                    w.to_detection_id,
                ]
            )
    _write_rows(wires_ws, wires_headers, wires_rows)

    assum_headers = ["Plan", "key", "value"]
    assum_rows: list[list[Any]] = []
    for b in bundles:
        label = _plan_label(b)
        for key, value in sorted(b.bom.assumptions.items(), key=lambda kv: str(kv[0])):
            assum_rows.append([label, str(key), value])
    _write_rows(assum_ws, assum_headers, assum_rows)

    corr_headers = [
        "Plan",
        "id",
        "entity",
        "entity_id",
        "action",
        "created_at",
        "before",
        "after",
    ]
    corr_rows: list[list[Any]] = []
    for b in bundles:
        label = _plan_label(b)
        for c in b.corrections:
            corr_rows.append(
                [
                    label,
                    c.id,
                    c.entity,
                    c.entity_id,
                    c.action,
                    c.created_at.isoformat() if c.created_at else "",
                    str(c.before) if c.before is not None else "",
                    str(c.after) if c.after is not None else "",
                ]
            )
    _write_rows(corr_ws, corr_headers, corr_rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "QuoteTitle",
            parent=base["Heading1"],
            fontSize=18,
            spaceAfter=8,
            textColor=colors.HexColor("#111827"),
        ),
        "h2": ParagraphStyle(
            "QuoteH2",
            parent=base["Heading2"],
            fontSize=12,
            spaceBefore=12,
            spaceAfter=6,
            textColor=colors.HexColor("#1F2937"),
        ),
        "body": ParagraphStyle(
            "QuoteBody",
            parent=base["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#374151"),
        ),
        "meta": ParagraphStyle(
            "QuoteMeta",
            parent=base["Normal"],
            fontSize=8,
            textColor=colors.HexColor("#6B7280"),
        ),
        "right": ParagraphStyle(
            "QuoteRight",
            parent=base["Normal"],
            fontSize=9,
            alignment=TA_RIGHT,
        ),
        "center": ParagraphStyle(
            "QuoteCenter",
            parent=base["Normal"],
            fontSize=9,
            alignment=TA_CENTER,
        ),
    }


def _table(data: list[list[Any]], col_widths: list[float] | None = None) -> Table:
    t = Table(data, colWidths=col_widths, hAlign="LEFT")
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
    ]
    if len(data) > 1:
        style.append(
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [colors.white, colors.HexColor("#F9FAFB")],
            )
        )
    t.setStyle(TableStyle(style))
    return t


def _money(value: float | None) -> str:
    if value is None:
        return "—"
    return f"${value:,.2f}"


def _quote_story_for_bundle(
    bundle: SheetExportBundle,
    styles: dict[str, ParagraphStyle],
    *,
    settings: Settings,
) -> list[Any]:
    story: list[Any] = []
    label = _plan_label(bundle)
    story.append(Paragraph("Electrical Material Quote", styles["title"]))
    story.append(
        Paragraph(
            f"<b>Project file:</b> {bundle.document.filename}<br/>"
            f"<b>Sheet:</b> {label} (page {bundle.sheet.page_index + 1})<br/>"
            f"<b>File hash:</b> {bundle.document.file_hash[:12]}…<br/>"
            f"<b>Date:</b> {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}<br/>"
            f"<b>Pipeline:</b> {bundle.pipeline_version}",
            styles["body"],
        )
    )
    story.append(Spacer(1, 0.15 * inch))

    # Summary counts (category totals)
    cat_totals = {k: v for k, v in bundle.takeoff.totals.items() if "." not in k}
    if cat_totals:
        story.append(Paragraph("Summary counts", styles["h2"]))
        rows = [["Category", "Qty"]] + [[k, v] for k, v in sorted(cat_totals.items())]
        story.append(_table(rows, col_widths=[3.5 * inch, 1.2 * inch]))

    story.append(Paragraph("Priced materials", styles["h2"]))
    bom_rows: list[list[Any]] = [
        ["Description", "Qty", "Unit", "Unit $", "Extended", "Source"]
    ]
    for line in bundle.bom.lines:
        bom_rows.append(
            [
                Paragraph(line.description or line.item_key, styles["body"]),
                f"{line.quantity:g}",
                line.unit,
                _money(line.unit_price),
                _money(line.extended_price),
                (line.price_source or "—")[:40],
            ]
        )
    bom_rows.append(
        ["TOTAL", "", "", "", _money(bundle.bom.total_usd), ""]
    )
    story.append(
        _table(
            bom_rows,
            col_widths=[2.4 * inch, 0.55 * inch, 0.45 * inch, 0.7 * inch, 0.8 * inch, 1.5 * inch],
        )
    )

    story.append(Paragraph("Assumptions &amp; price sources", styles["h2"]))
    assum_lines = [
        f"<b>{k}</b>: {v}"
        for k, v in sorted(bundle.bom.assumptions.items(), key=lambda kv: str(kv[0]))
    ]
    sources = sorted(
        {
            line.price_source
            for line in bundle.bom.lines
            if line.price_source
        }
    )
    if sources:
        assum_lines.append("<b>Price sources:</b> " + "; ".join(sources))
    story.append(
        Paragraph(
            "<br/>".join(assum_lines) if assum_lines else "No assumptions recorded.",
            styles["body"],
        )
    )

    if bundle.overlay_png and bundle.overlay_png.is_file():
        story.append(Paragraph("Plan viewport", styles["h2"]))
        img = Image(str(bundle.overlay_png))
        img._restrictSize(6.5 * inch, 3.2 * inch)  # noqa: SLF001
        story.append(img)

    story.append(Spacer(1, 0.2 * inch))
    story.append(
        Paragraph(
            "This quote is generated from the reviewed electrical takeoff. "
            "Counts and prices should be verified before purchasing.",
            styles["meta"],
        )
    )
    return story


def build_sheet_quote_pdf(bundle: SheetExportBundle, settings: Settings) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=0.65 * inch,
        rightMargin=0.65 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title=f"Quote — {_plan_label(bundle)}",
    )
    styles = _pdf_styles()
    doc.build(_quote_story_for_bundle(bundle, styles, settings=settings))
    return buf.getvalue()


def build_document_quote_pdf(bundles: list[SheetExportBundle], settings: Settings) -> bytes:
    buf = io.BytesIO()
    filename = bundles[0].document.filename if bundles else "document"
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=0.65 * inch,
        rightMargin=0.65 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title=f"Quote — {filename}",
    )
    styles = _pdf_styles()
    story: list[Any] = []
    for i, bundle in enumerate(bundles):
        if i:
            story.append(PageBreak())
        story.extend(_quote_story_for_bundle(bundle, styles, settings=settings))
    doc.build(story)
    return buf.getvalue()
